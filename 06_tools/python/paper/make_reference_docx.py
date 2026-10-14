"""Generate the docx style template (reference.docx) used by build_docx.py.

    /home/pi/miniconda3/envs/python-dev/bin/python 06_tools/python/paper/make_reference_docx.py \
        --out /tmp/reference.docx

The template starts from pandoc's own default reference.docx (so every style name pandoc writes exists) and changes:
  - page: A4, 25 mm margins, centered page number in the footer;
  - fonts: body NanumMyeongjo + Times New Roman, headings NanumGothic + Arial, code NanumGothicCoding (all slots).
    Every style and docDefaults get explicit ascii/hAnsi/eastAsia/cs fonts; theme font attributes (asciiTheme ...)
    are removed because they override explicit font names, and the theme itself is pointed at the same fonts;
  - sizes and spacing: body 10.5 pt justified, captions 9 pt left aligned (not italic; captions end with long
    source paths that would stretch justified lines), table text 9 pt, code 8 pt, references 9.5 pt with a
    hanging indent, headings and title black instead of pandoc's theme blue;
  - extra styles used by build_docx.py: "Table Text", "Table Heading", "Source Code" (no Hangul/Latin
    auto-spacing, so text diagrams keep their columns), "Horizontal Rule", "Footer";
  - a three-rule table style ("Table"): top and bottom rules, a thin rule under the header row;
  - settings: East Asian theme language ko-KR, Word 2013+ compatibility mode; font table lists the Korean
    fonts with their Korean names (Word matches either name).
The body holds a short sample of each style for previewing; pandoc ignores the body of a reference docx.
Output is byte-identical for the same pandoc and python-docx versions (fixed zip times and core dates).
Runs in conda env python-dev (python-docx, lxml, pandoc >= 3). Nothing else in the repository is read or written.
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

# ---------------------------------------------------------------------------------------------------------------
# Typography. Latin fonts follow the usual Korean journal pairing (Myeongjo + Times New Roman, Gothic + Arial);
# both Latin fonts ship with Word on Windows and macOS, and on Linux fontconfig maps them to the metric-compatible
# Liberation Serif / Liberation Sans, so line breaks seen here match Word. Code uses NanumGothicCoding in every
# slot: Latin is exactly half the width of Hangul, which keeps the manuscript's text diagrams aligned.
# ---------------------------------------------------------------------------------------------------------------
FONTS = {
    "body": {"latin": "Times New Roman", "east_asia": "NanumMyeongjo", "alt": "나눔명조"},
    "heading": {"latin": "Arial", "east_asia": "NanumGothic", "alt": "나눔고딕"},
    "mono": {"latin": "NanumGothicCoding", "east_asia": "NanumGothicCoding", "alt": "나눔고딕코딩"},
}
BODY_PT = 10.5
CAPTION_PT = 9.0
TABLE_PT = 9.0
CODE_PT = 8.0
FOOTER_PT = 9.0
PAGE_W_MM, PAGE_H_MM = 210, 297
MARGIN_MM = 25
HEADER_FOOTER_MM = 12.5
LANG_LATIN, LANG_EAST_ASIA = "en-US", "ko-KR"

FIXED_DATE = dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc)   # core.xml created/modified
ZIP_DATE = (1980, 1, 1, 0, 0, 0)                                 # earliest date a zip entry can hold

PANDOC_CANDIDATES = (
    Path("/home/pi/miniconda3/envs/python-dev/bin/pandoc"),
    Path(sys.prefix) / "bin" / "pandoc",
)
PANDOC_MIN_MAJOR = 3

THEME_ATTRS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
SAFE_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)

# style id -> (font family key, size pt or None, bold or None). Styles not listed get the body family.
HEADING_IDS = {"Title", "Subtitle", "Author", "AbstractTitle", "TOCHeading", "TableHeading"} | {
    f"Heading{i}" for i in range(1, 10)}
MONO_IDS = {"VerbatimChar", "SourceCode"}
# styles that by definition carry no formatting; Word expects them empty
SKIP_IDS = {"DefaultParagraphFont", "TableNormal", "NoList"}


class ToolError(RuntimeError):
    """Expected failure with a message for the user (missing tool, bad input)."""


# --------------------------------------------------------------------------- shared helpers (used by build_docx)

def find_pandoc(explicit: str | None = None) -> Path:
    """pandoc of conda env python-dev first, then of the running env, then PATH. Requires pandoc >= 3."""
    cands = [Path(explicit)] if explicit else list(PANDOC_CANDIDATES)
    if not explicit and shutil.which("pandoc"):
        cands.append(Path(shutil.which("pandoc")))
    for c in cands:
        if c.is_file() and os.access(c, os.X_OK):
            out = subprocess.run([str(c), "--version"], capture_output=True, text=True, check=True).stdout
            ver = out.split()[1]
            if int(ver.split(".")[0]) < PANDOC_MIN_MAJOR:
                raise ToolError(f"{c} is pandoc {ver}; pandoc >= {PANDOC_MIN_MAJOR} is required")
            return c
    tried = ", ".join(str(c) for c in cands) or "(none)"
    raise ToolError(f"pandoc not found (tried {tried}). Run in conda env python-dev or pass --pandoc PATH.")


def deterministic_zip(data: bytes) -> bytes:
    """Rewrite a docx zip with fixed entry times and attributes, [Content_Types].xml first, order otherwise kept."""
    src = zipfile.ZipFile(io.BytesIO(data))
    names = src.namelist()
    names.sort(key=lambda n: n != "[Content_Types].xml")    # stable: only moves the content-types entry
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as dst:
        for n in names:
            info = zipfile.ZipInfo(n, date_time=ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o644 << 16
            dst.writestr(info, src.read(n), compresslevel=9)
    return buf.getvalue()


def set_core_properties(document: docx.document.Document, title: str = "") -> None:
    cp = document.core_properties
    cp.title = title
    cp.author = ""
    cp.last_modified_by = ""
    cp.keywords = ""
    cp.revision = 1
    cp.created = FIXED_DATE.replace(tzinfo=None)
    cp.modified = FIXED_DATE.replace(tzinfo=None)


def set_rfonts(rpr: etree._Element, family: str) -> None:
    """Explicit fonts for every slot of an rPr; theme attributes are dropped (they would win over the names)."""
    f = FONTS[family]
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)              # rFonts is the first rPr child after rStyle
        rs = rpr.find(qn("w:rStyle"))
        if rs is not None:
            rpr.remove(rs)
            rpr.insert(0, rs)
    for a in THEME_ATTRS:
        rfonts.attrib.pop(qn(f"w:{a}"), None)
    rfonts.set(qn("w:ascii"), f["latin"])
    rfonts.set(qn("w:hAnsi"), f["latin"])
    rfonts.set(qn("w:cs"), f["latin"])
    rfonts.set(qn("w:eastAsia"), f["east_asia"])


def family_of(style_id: str) -> str:
    if style_id in HEADING_IDS or (style_id.endswith("Char") and style_id[:-4] in HEADING_IDS):
        return "heading"
    if style_id in MONO_IDS or style_id.endswith("Tok"):
        return "mono"
    return "body"


def strip_theme_fonts(root: etree._Element) -> int:
    n = 0
    for rf in root.iter(qn("w:rFonts")):
        for a in THEME_ATTRS:
            if rf.attrib.pop(qn(f"w:{a}"), None) is not None:
                n += 1
    return n


# --------------------------------------------------------------------------- style edits

def _get_or_add(parent: etree._Element, tag: str, successors: tuple[str, ...] = ()) -> etree._Element:
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        for s in successors:
            nxt = parent.find(qn(s))
            if nxt is not None:
                nxt.addprevious(el)
                break
        else:
            parent.append(el)
    return el


def _style(document, style_id: str, name: str, kind=WD_STYLE_TYPE.PARAGRAPH, based_on: str | None = None):
    """Existing style by id, or a new custom style."""
    for s in document.styles:
        if s.style_id == style_id:
            return s
    s = document.styles.add_style(name, kind)
    s.element.set(qn("w:styleId"), style_id)    # add_style derives the id from the name ('footer' -> 'footer')
    if based_on:
        s.base_style = _sid(document, based_on)
    s.quick_style = True
    return s


def _para(style, *, size=None, bold=None, italic=None, color=None, align=None, before=None, after=None,
          line=None, keep_next=None, keep_lines=None, left=None, first_line=None) -> None:
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    pf = getattr(style, "paragraph_format", None)
    if pf is None:
        return
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if keep_lines is not None:
        pf.keep_together = keep_lines
    if left is not None:
        pf.left_indent = Mm(left)
    if first_line is not None:
        pf.first_line_indent = Mm(first_line)


def _drop_color_theme(style) -> None:
    """Heading colors in pandoc's template are theme colors (blue); make them plain black."""
    rpr = style.element.find(qn("w:rPr"))
    if rpr is not None:
        c = rpr.find(qn("w:color"))
        if c is not None:
            rpr.remove(c)


def _border(tag: str, sz: int, color: str = "000000", space: int = 0) -> etree._Element:
    b = OxmlElement(tag)
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    return b


PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                  "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                  "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr")
PPR_AFTER_SHD = PPR_AFTER_PBDR[1:]


def _ppr(style) -> etree._Element:
    return style.element.get_or_add_pPr()


def edit_styles(document) -> None:
    st = document.styles
    root = st.element

    # docDefaults: body fonts, body size, Korean as the East Asian language
    rpr_default = root.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    set_rfonts(rpr_default, "body")
    rpr_default.find(qn("w:sz")).set(qn("w:val"), str(int(BODY_PT * 2)))
    rpr_default.find(qn("w:szCs")).set(qn("w:val"), str(int(BODY_PT * 2)))
    lang = rpr_default.find(qn("w:lang"))
    lang.set(qn("w:val"), LANG_LATIN)
    lang.set(qn("w:eastAsia"), LANG_EAST_ASIA)
    ppr_default = root.find(qn("w:docDefaults")).find(qn("w:pPrDefault")).find(qn("w:pPr"))
    sp = ppr_default.find(qn("w:spacing"))
    sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), "276")          # 1.15 lines
    sp.set(qn("w:lineRule"), "auto")

    # body text
    for sid in ("BodyText", "FirstParagraph"):
        _para(_sid(document, sid), align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6)
    _para(_sid(document, "Compact"), before=1.8, after=1.8)
    _para(_sid(document, "BlockText"), size=9.5, before=3, after=6, left=6, first_line=0)
    bt = _ppr(_sid(document, "BlockText"))
    ind = bt.find(qn("w:ind"))
    for a in ("w:right",):
        ind.attrib.pop(qn(a), None)
    pbdr = _get_or_add(bt, "w:pBdr", PPR_AFTER_PBDR)
    pbdr.append(_border("w:left", 12, "A6A6A6", 6))

    # title block and headings
    _para(_sid(document, "Title"), size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6,
          line=1.0)
    _para(_sid(document, "Subtitle"), size=13, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=12)
    for sid in ("Title", "Subtitle", "TitleChar", "SubtitleChar"):
        _drop_color_theme(_sid(document, sid))
    heading = {1: (14, 18, 6), 2: (12, 12, 6), 3: (11, 10, 4)}
    for lvl in range(1, 10):
        size, before, after = heading.get(lvl, (BODY_PT, 8, 3))
        for sid in (f"Heading{lvl}", f"Heading{lvl}Char"):
            s = _sid(document, sid)
            _drop_color_theme(s)
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = False
        _para(_sid(document, f"Heading{lvl}"), before=before, after=after, keep_next=True, keep_lines=True)

    # captions and figures
    _para(_sid(document, "Caption"), size=CAPTION_PT, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=3,
          after=10)
    _para(_sid(document, "ImageCaption"), keep_next=False)
    _para(_sid(document, "TableCaption"), before=6, after=3, keep_next=True)
    _para(_sid(document, "Figure"), align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3, keep_next=True,
          keep_lines=True, line=1.0)
    _para(_sid(document, "CaptionedFigure"), keep_next=True)

    # tables: paragraph styles for cells (set by build_docx.py) and the three-rule table style
    tt = _style(document, "TableText", "Table Text", based_on="Normal")
    _para(tt, size=TABLE_PT, align=WD_ALIGN_PARAGRAPH.LEFT, before=1, after=1, line=1.0)
    th = _style(document, "TableHeading", "Table Heading", based_on="TableText")
    _para(th, bold=True, keep_next=True)
    edit_table_style(_sid(document, "Table"))

    # code
    sc = _style(document, "SourceCode", "Source Code", based_on="Normal")
    _para(sc, size=CODE_PT, align=WD_ALIGN_PARAGRAPH.LEFT, before=3, after=6, line=1.0, keep_lines=True)
    sc_ppr = _ppr(sc)
    shd = _get_or_add(sc_ppr, "w:shd", PPR_AFTER_SHD)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F3F3F3")
    wrap = _get_or_add(sc_ppr, "w:wordWrap", PPR_AFTER_PBDR[5:])
    wrap.set(qn("w:val"), "0")           # long code lines may break anywhere instead of overflowing
    for tag, i in (("w:autoSpaceDE", 8), ("w:autoSpaceDN", 9)):   # no extra gap at Hangul/Latin boundaries,
        _get_or_add(sc_ppr, tag, PPR_AFTER_PBDR[i:]).set(qn("w:val"), "0")   # which would break diagrams
    vc = _sid(document, "VerbatimChar")
    vrpr = vc.element.find(qn("w:rPr"))
    for tag in ("w:sz", "w:szCs"):       # inherit the paragraph size (body, table cell or code block)
        el = vrpr.find(qn(tag))
        if el is not None:
            vrpr.remove(el)

    # references (build_docx.py gives this style to the paragraphs under a "참고문헌"/"References" heading)
    bib = _sid(document, "Bibliography")
    _para(bib, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=4, left=8, first_line=-8)

    # horizontal rule (build_docx.py replaces pandoc's VML rule with an empty paragraph of this style)
    hr = _style(document, "HorizontalRule", "Horizontal Rule", based_on="Normal")
    _para(hr, size=4, before=6, after=12, line=1.0)
    _get_or_add(_ppr(hr), "w:pBdr", PPR_AFTER_PBDR).append(_border("w:bottom", 4, "808080", 1))

    # misc
    _para(_sid(document, "FootnoteText"), size=9)
    hl = _sid(document, "Hyperlink")
    _drop_color_theme(hl)
    hl.font.color.rgb = RGBColor.from_string("1F3A93")
    ft = _style(document, "Footer", "footer", based_on="Normal")
    _para(ft, size=FOOTER_PT, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0)

    # explicit fonts on every style (after the new styles exist)
    for s in root.iter(qn("w:style")):
        sid = s.get(qn("w:styleId"))
        if sid in SKIP_IDS or s.get(qn("w:type")) == "numbering":
            continue
        rpr = s.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            # rPr goes after pPr, before tblPr/trPr/tcPr/tblStylePr
            nxt = next((c for c in s if c.tag in {qn("w:tblPr"), qn("w:trPr"), qn("w:tcPr"),
                                                  qn("w:tblStylePr")}), None)
            if nxt is not None:
                nxt.addprevious(rpr)
            else:
                s.append(rpr)
        set_rfonts(rpr, family_of(sid))
    strip_theme_fonts(root)


def _sid(document, style_id: str):
    """Style by id. python-docx's lookup by name maps built-in names to lower-case aliases, so ids are safer."""
    for s in document.styles:
        if s.style_id == style_id:
            return s
    raise KeyError(style_id)


def edit_table_style(style) -> None:
    """Three-rule table: 1 pt top and bottom, 0.5 pt under the header row, centered, small cell margins."""
    el = style.element
    tblpr = el.find(qn("w:tblPr"))
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblpr.find(qn("w:tblInd")).addprevious(jc)
    borders = OxmlElement("w:tblBorders")
    borders.append(_border("w:top", 8))
    borders.append(_border("w:bottom", 8))
    tblpr.find(qn("w:tblInd")).addnext(borders)
    mar = tblpr.find(qn("w:tblCellMar"))
    for side, w in (("top", 28), ("bottom", 28), ("left", 85), ("right", 85)):
        mar.find(qn(f"w:{side}")).set(qn("w:w"), str(w))
    first = el.find(qn("w:tblStylePr"))
    b = first.find(qn("w:tcPr")).find(qn("w:tcBorders")).find(qn("w:bottom"))
    for k, v in (("sz", "4"), ("space", "0"), ("color", "000000")):
        b.set(qn(f"w:{k}"), v)
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:b"))
    first.find(qn("w:tcPr")).addprevious(rpr)
    first.find(qn("w:tcPr")).addprevious(OxmlElement("w:trPr"))
    first.find(qn("w:trPr")).append(OxmlElement("w:tblHeader"))


# --------------------------------------------------------------------------- page, theme, fonts, settings

def edit_section(document) -> None:
    sec = document.sections[0]
    sec.page_width, sec.page_height = Mm(PAGE_W_MM), Mm(PAGE_H_MM)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(MARGIN_MM)
    sec.header_distance = sec.footer_distance = Mm(HEADER_FOOTER_MM)
    sec.gutter = 0
    footer = sec.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.style = _sid(document, "Footer")
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = p.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text == " PAGE ":
            it = OxmlElement("w:instrText")
            it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text


def _part(document, partname: str):
    for p in document.part.package.iter_parts():
        if str(p.partname) == partname:
            return p
    raise KeyError(partname)


def edit_theme(document) -> None:
    part = _part(document, "/word/theme/theme1.xml")
    root = etree.fromstring(part.blob, SAFE_PARSER)
    for which, fam in (("majorFont", "heading"), ("minorFont", "body")):
        fe = root.find(f".//{{{NS_A}}}{which}")
        latin = fe.find(f"{{{NS_A}}}latin")
        latin.set("typeface", FONTS[fam]["latin"])
        latin.attrib.pop("panose", None)
        for f in fe.findall(f"{{{NS_A}}}font"):
            if f.get("script") == "Hang":
                f.set("typeface", FONTS[fam]["east_asia"])
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def edit_font_table(document) -> None:
    """Declare the Korean fonts with the Hangul charset and their Korean names, so Word can match either name."""
    part = _part(document, "/word/fontTable.xml")
    root = etree.fromstring(part.blob, SAFE_PARSER)
    have = {f.get(qn("w:name")) for f in root.findall(qn("w:font"))}
    for fam in ("body", "heading", "mono"):
        f = FONTS[fam]
        for name, korean in ((f["east_asia"], True), (f["latin"], f["latin"] == f["east_asia"])):
            if name in have:
                continue
            have.add(name)
            el = etree.SubElement(root, qn("w:font"))
            el.set(qn("w:name"), name)
            if korean:
                etree.SubElement(el, qn("w:altName")).set(qn("w:val"), f["alt"])
            etree.SubElement(el, qn("w:charset")).set(qn("w:val"), "81" if korean else "00")
            etree.SubElement(el, qn("w:family")).set(
                qn("w:val"), "modern" if fam == "mono" else ("roman" if fam == "body" else "swiss"))
            etree.SubElement(el, qn("w:pitch")).set(qn("w:val"), "fixed" if fam == "mono" else "variable")
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def edit_settings(document) -> None:
    root = document.settings.element
    tfl = root.find(qn("w:themeFontLang"))
    tfl.set(qn("w:val"), LANG_LATIN)
    tfl.set(qn("w:eastAsia"), LANG_EAST_ASIA)
    if root.find(qn("w:compat")) is None:     # open in Word 2013+ layout mode, not compatibility mode
        compat = OxmlElement("w:compat")
        cs = OxmlElement("w:compatSetting")
        for k, v in (("name", "compatibilityMode"), ("uri", "http://schemas.microsoft.com/office/word"),
                     ("val", "15")):
            cs.set(qn(f"w:{k}"), v)
        compat.append(cs)
        root.find(qn("w:footnotePr")).addnext(compat)


# --------------------------------------------------------------------------- preview body

def write_sample_body(document) -> None:
    body = document.element.body
    sect = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect:
            body.remove(child)
    d = document

    def add(text: str, sid: str):
        return d.add_paragraph(text, style=_sid(d, sid))

    add("논문 제목 견본 (Title)", "Title")
    add("— 부제 견본 (Subtitle)", "Subtitle")
    add("1. 장 제목 (Heading 1)", "Heading1")
    add("1.1 절 제목 (Heading 2)", "Heading2")
    add("1.1.1 항 제목 (Heading 3)", "Heading3")
    p = add("본문 첫 문단은 나눔명조 10.5 pt, 라틴 문자는 Times New Roman이다. "
            "The quick brown fox jumps over the lazy dog 0123456789. ", "FirstParagraph")
    p.add_run("inline_code()").style = _sid(d, "VerbatimChar")
    add("둘째 문단부터는 Body Text 스타일이다. 양쪽 정렬, 문단 뒤 6 pt.", "BodyText")
    add("인용 문단(Block Text)은 9.5 pt이고 왼쪽에 회색 세로줄이 있다.", "BlockText")
    add("[그림 자리 — Figure 스타일, 가운데 정렬]", "Figure")
    add("그림 1. 그림 캡션(Image Caption)은 9 pt, 왼쪽 정렬, 기울이지 않는다.", "ImageCaption")
    add("표 1. 표 캡션(Table Caption)은 표 위에 둔다.", "TableCaption")
    t = d.add_table(rows=3, cols=3)
    t.style = _sid(d, "Table")
    for r, row in enumerate([("항목", "값", "단위"), ("CPU", "1.842", "ms/frame"), ("지연", "101.20", "ms")]):
        for c, text in enumerate(row):
            cell = t.cell(r, c)
            cell.paragraphs[0].text = text
            cell.paragraphs[0].style = _sid(d, "TableHeading" if r == 0 else "TableText")
    add("free ──acquire──▶ 전처리 ──▶ free\n  ▼ ① 코드 블록(Source Code)", "SourceCode")
    add("", "HorizontalRule")
    add("[R1] 참고문헌 견본(Bibliography). 왼쪽 정렬, 9.5 pt, 둘째 줄부터 8 mm 내어쓰기. "
        "Available: https://example.org/a/very/long/url/that/cannot/break", "Bibliography")


# --------------------------------------------------------------------------- main

def pandoc_default_reference(pandoc: Path) -> bytes:
    return subprocess.run([str(pandoc), "--print-default-data-file", "reference.docx"],
                          capture_output=True, check=True).stdout


def build_reference(out: Path, pandoc: Path) -> Path:
    document = docx.Document(io.BytesIO(pandoc_default_reference(pandoc)))
    edit_styles(document)          # first: the footer paragraph needs the Footer style
    edit_section(document)
    edit_theme(document)
    edit_font_table(document)
    edit_settings(document)
    write_sample_body(document)
    set_core_properties(document)
    buf = io.BytesIO()
    document.save(buf)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(deterministic_zip(buf.getvalue()))
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Generate the docx style template (reference.docx) for build_docx.py.")
    ap.add_argument("--out", required=True, type=Path, help="output .docx path")
    ap.add_argument("--pandoc", default=None, help="pandoc executable (default: python-dev env pandoc, then PATH)")
    a = ap.parse_args(argv)
    try:
        pandoc = find_pandoc(a.pandoc)
        out = build_reference(a.out.resolve(), pandoc)
    except ToolError as e:
        print(f"error: {e}", file=sys.stderr)
        return 2
    print(f"wrote: {out}")
    print(f"pandoc: {pandoc}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
